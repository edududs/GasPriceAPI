"""The INMETRO light vehicle label table (PBEV): km/L in city and highway, for ethanol and for gasoline
or diesel, per brand, model and version.

The workbook's header spans several rows (merged "Etanol" over "Cidade | Estrada"), so each column's
name is built from every header row above its data, with merged group titles carried to the right.
"""

import re
from collections.abc import Callable, Sequence
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import cast
from zipfile import BadZipFile

from openpyxl import load_workbook
from pydantic import ValidationError

from gasprice.prices.domain import DomainError, Fuel, normalize, parse_decimal
from gasprice.trips.application import VehicleHarvest, VehicleSourceError
from gasprice.trips.domain import Efficiency, Vehicle, VehicleSource

type Cell = object
type CellReader = Callable[[str], Cell]
HEADER_SEARCH_ROWS = 40
SUBHEADER_ROWS = 3


class PbevWorkbook:
    source = VehicleSource.PBEV

    def __init__(self, path: Path, *, year: int | None = None) -> None:
        self._path = path
        self._year = year

    def read(self) -> VehicleHarvest:
        try:
            data = self._path.read_bytes()
        except OSError as error:
            msg = f"não foi possível ler {self._path}: {error}"
            raise VehicleSourceError(msg) from error
        return parse_pbev(data, year=self._year)


def parse_pbev(data: bytes, *, year: int | None = None) -> VehicleHarvest:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, OSError, KeyError, ValueError) as error:
        msg = f"o arquivo não é uma planilha .xlsx válida: {error}"
        raise VehicleSourceError(msg) from error
    try:
        for sheet in workbook.worksheets:
            rows = [tuple(cast("Sequence[Cell]", row)) for row in sheet.iter_rows(values_only=True)]
            columns = _find_columns(rows)
            if columns is not None:
                return _harvest(rows, *columns, year=year)
    finally:
        workbook.close()
    msg = "nenhuma aba com MARCA, MODELO e km/l de cidade e estrada"
    raise VehicleSourceError(msg)


def _text(cell: Cell) -> str:
    return normalize(str(cell)) if cell is not None else ""


def _find_columns(rows: Sequence[Sequence[Cell]]) -> tuple[int, dict[str, int]] | None:
    """The header row index and the position of each field, or None if this sheet has no table."""
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        names = [_text(cell) for cell in row]
        if "MARCA" not in names or "MODELO" not in names:
            continue
        block = [list(rows[i]) for i in range(index, min(index + SUBHEADER_ROWS, len(rows)))]
        composite = _composite_names(block)
        columns = _match(composite)
        if columns is not None:
            return index, columns
    return None


def _composite_names(block: Sequence[Sequence[Cell]]) -> list[str]:
    width = max(len(row) for row in block)
    names = [""] * width
    for depth, row in enumerate(block):
        carried = ""
        for column in range(width):
            text = _text(row[column]) if column < len(row) else ""
            below = (
                block[depth + 1][column]
                if depth + 1 < len(block) and column < len(block[depth + 1])
                else None
            )
            if text:
                carried = text
            elif below is not None and carried:
                text = carried  # a merged group title spans this column too
            else:
                carried = ""
            names[column] = f"{names[column]} {text}".strip()
    return names


def _match(names: Sequence[str]) -> dict[str, int] | None:
    def find(*tokens: str, exclude: tuple[str, ...] = ()) -> int | None:
        for position, name in enumerate(names):
            words = set(re.split(r"[^A-Z0-9]+", name))
            if all(token in words for token in tokens) and not words & set(exclude):
                return position
        return None

    found = {
        "brand": find("MARCA"),
        "model": find("MODELO"),
        "version": find("VERSAO"),
        "year": find("ANO"),
        "fuel": find("COMBUSTIVEL", exclude=("CIDADE", "ESTRADA")),
        "ethanol_city": find("ETANOL", "CIDADE"),
        "ethanol_highway": find("ETANOL", "ESTRADA"),
        # The other pair is titled "Gasolina ou Diesel" in the label table; a plain sheet may say either.
        "other_city": find("GASOLINA", "CIDADE") or find("DIESEL", "CIDADE"),
        "other_highway": find("GASOLINA", "ESTRADA") or find("DIESEL", "ESTRADA"),
    }
    columns = {field: position for field, position in found.items() if position is not None}
    pairs = ({"other_city", "other_highway"}, {"ethanol_city", "ethanol_highway"})
    has_pair = any(pair <= columns.keys() for pair in pairs)
    return columns if {"brand", "model"} <= columns.keys() and has_pair else None


def _harvest(
    rows: Sequence[Sequence[Cell]], header: int, columns: dict[str, int], *, year: int | None
) -> VehicleHarvest:
    vehicles: list[Vehicle] = []
    skipped: list[str] = []
    for index in range(header + 1, len(rows)):
        row = rows[index]

        def cell(field: str, row: Sequence[Cell] = row) -> Cell:
            position = columns.get(field)
            return row[position] if position is not None and position < len(row) else None

        brand, model = _clean(cell("brand")), _clean(cell("model"))
        if not brand or not model:
            continue  # sub-header, spacer or footnote
        try:
            efficiencies = _efficiencies(cell)
            if not efficiencies:
                skipped.append(f"linha {index + 1}: sem km/l de cidade e estrada")
                continue
            vehicles.append(
                Vehicle(
                    id="",
                    source=VehicleSource.PBEV,
                    brand=brand.title() if brand.isupper() else brand,
                    model=model,
                    version=_clean(cell("version")),
                    year=_year(cell("year")) or year,
                    efficiencies=efficiencies,
                )
            )
        except (DomainError, ValidationError, ValueError) as error:
            reason = error.errors()[0]["msg"] if isinstance(error, ValidationError) else str(error)
            skipped.append(f"linha {index + 1}: {reason}")
    return VehicleHarvest(vehicles=tuple(vehicles), skipped=tuple(skipped))


def _efficiencies(cell: CellReader) -> dict[Fuel, Efficiency]:
    efficiencies: dict[Fuel, Efficiency] = {}
    ethanol = _pair(cell("ethanol_city"), cell("ethanol_highway"))
    if ethanol is not None:
        efficiencies[Fuel.ETHANOL] = ethanol
    other = _pair(cell("other_city"), cell("other_highway"))
    if other is not None:
        efficiencies[Fuel.DIESEL_S10 if "D" in _fuel_codes(cell("fuel")) else Fuel.GASOLINE] = other
    return efficiencies


def _fuel_codes(value: Cell) -> set[str]:
    """`F` flex, `G` gasoline, `D` diesel, `E` ethanol, or the words themselves."""
    text = _text(value)
    if "DIESEL" in text or text == "D":
        return {"D"}
    return {text[:1]} if text else set()


def _pair(city: Cell, highway: Cell) -> Efficiency | None:
    city_value, highway_value = _number(city), _number(highway)
    if city_value is None or highway_value is None:
        return None
    return Efficiency(city=city_value, highway=highway_value)


def _number(value: Cell) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return Decimal(str(value)) if value > 0 else None
    text = str(value).strip()
    if not text or text in {"-", "*", "N/A", "NA"}:
        return None
    return parse_decimal(text)


def _clean(value: Cell) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _year(value: Cell) -> int | None:
    text = _clean(value)
    digits = "".join(char for char in text if char.isdigit())
    return int(digits[:4]) if len(digits) >= 4 else None  # noqa: PLR2004
