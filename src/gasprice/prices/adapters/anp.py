"""The ANP weekly fuel price survey (Levantamento de Preços de Combustíveis, LPC).

The ANP publishes each week a workbook whose `ESTADOS` sheet has one row per state and product:
survey dates, number of stations, and the mean, minimum and maximum pump prices. The historical
workbook since 2013 has the same columns. Neither has a fixed header row, so the parser looks for it.
"""

import re
from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import cast
from urllib.parse import urljoin
from zipfile import BadZipFile

import httpx
from openpyxl import load_workbook
from pydantic import ValidationError

from gasprice.prices.adapters.http import get_bytes
from gasprice.prices.application import Harvest, SourceUnavailableError
from gasprice.prices.domain import (
    DomainError,
    Fuel,
    InvalidPriceError,
    PriceReport,
    Source,
    State,
    normalize,
    parse_decimal,
)

type Cell = object
HEADER_SEARCH_ROWS = 40

COLUMNS = {
    "period_start": "DATA INICIAL",
    "period_end": "DATA FINAL",
    "state": "ESTADO",
    "product": "PRODUTO",
    "stations": "NUMERO DE POSTOS PESQUISADOS",
    "average": "PRECO MEDIO REVENDA",
    "minimum": "PRECO MINIMO REVENDA",
    "maximum": "PRECO MAXIMO REVENDA",
}
REQUIRED = ("period_start", "period_end", "state", "product", "average")
EXACT = frozenset({"state", "product"})
"""`ESTADO` must not match `ESTADOS`. The price columns carry suffixes in some years: they match by prefix."""

_WORKBOOK_LINK = re.compile(r"""href=["']([^"']+\.xlsx)["']""", re.IGNORECASE)
_DATE_IN_NAME = re.compile(r"(\d{4})[-_](\d{2})[-_](\d{2})")


@dataclass(frozen=True)
class _Header:
    row: int
    columns: dict[str, int]


class AnpSource:
    """Read the survey from a workbook URL or file; with neither, find this week's on the ANP page."""

    source = Source.ANP

    def __init__(self, client: httpx.Client, *, page_url: str, workbook: str | None = None) -> None:
        self._client = client
        self._page_url = page_url
        self._workbook = workbook

    def fetch(self) -> Harvest:
        return parse_workbook(self._read_workbook())

    def _read_workbook(self) -> bytes:
        if self._workbook is None:
            page = get_bytes(self._client, self._page_url).decode("utf-8", errors="replace")
            return get_bytes(self._client, find_latest_workbook(page, self._page_url))
        if self._workbook.startswith(("http://", "https://")):
            return get_bytes(self._client, self._workbook)
        path = Path(self._workbook)
        try:
            return path.read_bytes()
        except OSError as error:
            msg = f"não foi possível ler {path}: {error}"
            raise SourceUnavailableError(msg) from error


def find_latest_workbook(page: str, base_url: str) -> str:
    """The weekly workbook linked from the survey page with the latest date in its name."""
    links = [urljoin(base_url, link) for link in _WORKBOOK_LINK.findall(page)]
    weekly = [link for link in links if "semanal" in link.lower()] or links
    if not weekly:
        msg = f"nenhuma planilha .xlsx encontrada em {base_url}"
        raise SourceUnavailableError(msg)
    return max(weekly, key=_last_date_in_name)


def _last_date_in_name(link: str) -> str:
    dates = _DATE_IN_NAME.findall(link.rsplit("/", 1)[-1])
    return "-".join(dates[-1]) if dates else ""


def parse_workbook(data: bytes) -> Harvest:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, OSError, KeyError, ValueError) as error:
        msg = f"o arquivo não é uma planilha .xlsx válida: {error}"
        raise SourceUnavailableError(msg) from error
    try:
        for sheet in workbook.worksheets:
            rows = [tuple(cast("Sequence[Cell]", row)) for row in sheet.iter_rows(values_only=True)]
            header = _find_header(rows)
            if header is not None:
                return _harvest(rows, header)
    finally:
        workbook.close()
    msg = "nenhuma aba com as colunas de preço por estado (ESTADO, PRODUTO, PREÇO MÉDIO REVENDA)"
    raise SourceUnavailableError(msg)


def _find_header(rows: Sequence[Sequence[Cell]]) -> _Header | None:
    for index, row in enumerate(rows[:HEADER_SEARCH_ROWS]):
        names = [normalize(str(cell)) if cell is not None else "" for cell in row]
        if any(name.startswith("MUNICIPIO") for name in names):
            return None  # the municipalities sheet repeats ESTADO and PRODUTO; it is not this table
        columns: dict[str, int] = {}
        for field, title in COLUMNS.items():
            for position, name in enumerate(names):
                if name == title if field in EXACT else name.startswith(title):
                    columns[field] = position
                    break
        if all(field in columns for field in REQUIRED):
            return _Header(row=index, columns=columns)
    return None


def _harvest(rows: Sequence[Sequence[Cell]], header: _Header) -> Harvest:
    reports: list[PriceReport] = []
    skipped: list[str] = []
    for number, row in _data_rows(rows, header):
        try:
            reports.append(_report(row, header.columns))
        except (DomainError, ValidationError, ValueError) as error:
            reason = error.errors()[0]["msg"] if isinstance(error, ValidationError) else str(error)
            skipped.append(f"linha {number}: {reason}")
    return Harvest(reports=tuple(reports), skipped=tuple(skipped))


def _data_rows(rows: Sequence[Sequence[Cell]], header: _Header) -> Iterator[tuple[int, Sequence[Cell]]]:
    product = header.columns["product"]
    average = header.columns["average"]
    for index in range(header.row + 1, len(rows)):
        row = rows[index]
        if _blank(_at(row, product)) and _blank(_at(row, average)):
            continue  # spacer or footnote
        yield index + 1, row


def _report(row: Sequence[Cell], columns: dict[str, int]) -> PriceReport:
    def cell(field: str) -> Cell:
        return _at(row, columns[field]) if field in columns else None

    return PriceReport(
        source=Source.ANP,
        state=State.from_name(str(cell("state"))),
        fuel=Fuel.from_product_name(str(cell("product"))),
        period_start=_date(cell("period_start")),
        period_end=_date(cell("period_end")),
        average=_price(cell("average")),
        minimum=_optional_price(cell("minimum")),
        maximum=_optional_price(cell("maximum")),
        stations=_optional_int(cell("stations")),
    )


def _at(row: Sequence[Cell], position: int) -> Cell:
    return row[position] if position < len(row) else None


def _blank(value: Cell) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _date(value: Cell) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for pattern in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
        try:
            return datetime.strptime(text, pattern).date()  # noqa: DTZ007  a calendar date, no instant
        except ValueError:
            continue
    msg = f"data inválida: {value!r}"
    raise ValueError(msg)


def _price(value: Cell) -> Decimal:
    if isinstance(value, bool):
        msg = f"não é um preço: {value!r}"
        raise InvalidPriceError(msg)
    if isinstance(value, int | float):
        return Decimal(str(value))
    return parse_decimal(str(value))


def _optional_price(value: Cell) -> Decimal | None:
    if _blank(value) or (isinstance(value, str) and value.strip() == "-"):
        return None
    return _price(value)


def _optional_int(value: Cell) -> int | None:
    if _blank(value):
        return None
    if isinstance(value, int | float):
        return int(value)
    return int(str(value).strip().replace(".", ""))
